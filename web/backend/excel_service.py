from __future__ import annotations

from pathlib import Path
from openpyxl import load_workbook

from .models import Recipient

REQUIRED_COLUMNS = {"store_code", "store_name", "email"}

def load_recipients(path: str | Path) -> list[Recipient]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().lower() for value in next(rows)]
    index = {name: i for i, name in enumerate(headers)}
    missing = REQUIRED_COLUMNS - index.keys()
    if missing:
        raise ValueError(f"Missing columns: {', '.join(sorted(missing))}")

    recipients = []
    for row in rows:
        email = str(row[index["email"]] or "").strip()
        if not email:
            continue
        recipients.append(
            Recipient(
                store_code=str(row[index["store_code"]] or "").strip(),
                store_name=str(row[index["store_name"]] or "").strip(),
                email=email,
                region=str(row[index["region"]] or "").strip() if "region" in index else "",
            )
        )
    return recipients
